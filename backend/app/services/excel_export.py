from __future__ import annotations

import io
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, PieChart, RadarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from .. import models, schemas
from . import aggregations


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def _write_df(ws, df: pd.DataFrame, start_row: int = 1, title: str | None = None) -> int:
    """Write a dataframe to worksheet starting at start_row. Returns next free row."""
    row = start_row
    if title:
        ws.cell(row=row, column=1, value=title).font = Font(bold=True, size=14)
        row += 1
    if df.empty:
        ws.cell(row=row, column=1, value="(no data)")
        return row + 2

    cols = list(df.columns)
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=c, value=str(name))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    row += 1

    for _, rec in df.iterrows():
        for c, name in enumerate(cols, start=1):
            v = rec[name]
            if hasattr(v, "to_pydatetime"):
                v = v.to_pydatetime()
            ws.cell(row=row, column=c, value=None if pd.isna(v) else v)
        row += 1

    for c, name in enumerate(cols, start=1):
        if len(df):
            lens = [len(str(v)) for v in df[name].tolist()]
            max_len = max(lens) if lens else 12
        else:
            max_len = 12
        ws.column_dimensions[get_column_letter(c)].width = max(12, min(40, max_len + 2))
    return row + 1


def _write_pivot(ws, df: pd.DataFrame, start_row: int, title: str) -> int:
    ws.cell(row=start_row, column=1, value=title).font = Font(bold=True, size=14)
    if df.empty:
        ws.cell(row=start_row + 1, column=1, value="(no data)")
        return start_row + 3
    flat = df.copy()
    flat.columns = [" / ".join(str(x) for x in col) if isinstance(col, tuple) else str(col) for col in flat.columns]
    flat = flat.reset_index()
    return _write_df(ws, flat, start_row + 1)


# ---------- native chart helpers ----------

def _analytics_points_to_table(points: list[dict]) -> pd.DataFrame:
    """Turn the /analytics/query `points` array into a wide DataFrame.

    Columns: [X, <series 1>, <series 2>, ...]. One row per unique X value,
    preserving first-seen ordering. Missing series/x pairs are filled with 0.
    """
    if not points:
        return pd.DataFrame()
    seen_x: list[str] = []
    seen_s: list[str] = []
    cells: dict[tuple[str, str], float] = {}
    for p in points:
        x = str(p["x"])
        s = str(p.get("series") or "value")
        if x not in seen_x:
            seen_x.append(x)
        if s not in seen_s:
            seen_s.append(s)
        cells[(x, s)] = float(p["y"])
    data = {"X": seen_x}
    for s in seen_s:
        data[s] = [cells.get((x, s), 0) for x in seen_x]
    return pd.DataFrame(data)


def _build_native_chart(chart_type: str, title: str, data_rows: int, data_cols: int, ws, data_start_row: int):
    """Given a data table starting at `data_start_row` (header row), build the
    appropriate openpyxl chart object referencing that range.

    The table looks like:
      data_start_row:        X | series1 | series2 | ...
      data_start_row+1..N:   <val> | <y1> | <y2> | ...
    """
    cls_map = {
        "bar": BarChart,
        "line": LineChart,
        "pie": PieChart,
        "radar": RadarChart,
    }
    chart_cls = cls_map.get(chart_type, BarChart)
    chart = chart_cls()
    if chart_type == "radar":
        chart.type = "filled"
    chart.title = title
    chart.style = 10

    # Series (exclude the X column, so start at column 2)
    values_ref = Reference(
        ws,
        min_col=2,
        max_col=data_cols,
        min_row=data_start_row,
        max_row=data_start_row + data_rows,
    )
    cats_ref = Reference(
        ws,
        min_col=1,
        max_col=1,
        min_row=data_start_row + 1,
        max_row=data_start_row + data_rows,
    )
    chart.add_data(values_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    
    if hasattr(chart, "x_axis") and chart.x_axis is not None:
        chart.x_axis.delete = False
    if hasattr(chart, "y_axis") and chart.y_axis is not None:
        chart.y_axis.delete = False
        
    chart.height = 10
    chart.width = 20
    return chart


def _build_occupation_chart(title: str, data_rows: int, ws, data_start_row: int):
    """Specific helper for the occupation breakdown table.
    Table: Category | Question | 5 | 4 | 3 | 2 | 1
    Questions are in column 2, Values are in columns 3..7
    """
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "clustered"
    chart.title = title
    chart.style = 10
    chart.dLbls = DataLabelList(showVal=True)
    chart.x_axis.delete = False
    chart.y_axis.delete = False

    # Data: columns 3 to 7 (scores 5 down to 1)
    values_ref = Reference(
        ws,
        min_col=3,
        max_col=7,
        min_row=data_start_row,
        max_row=data_start_row + data_rows
    )
    # X axis labels: column 2 (Question text)
    cats_ref = Reference(
        ws,
        min_col=2,
        max_col=2,
        min_row=data_start_row + 1,
        max_row=data_start_row + data_rows
    )
    chart.add_data(values_ref, titles_from_data=True)
    chart.set_categories(cats_ref)

    # Standard Likert colors: 5=Green, 4=Lime, 3=Yellow, 2=Orange, 1=Red
    # Series index in Excel starts from 0. Our cols are 5, 4, 3, 2, 1.
    colors = ["22c55e", "84cc16", "eab308", "f97316", "ef4444"]
    for i, color in enumerate(colors):
        if i < len(chart.series):
            chart.series[i].graphicalProperties.solidFill = color

    chart.height = 10
    chart.width = 20
    return chart


def _occupation_breakdown_to_table(points: list[dict], survey: models.Survey) -> pd.DataFrame:
    """Matrix: [Category, Question, 5, 4, 3, 2, 1]"""
    if not points:
        return pd.DataFrame()

    matrix = {}
    for p in points:
        q_text = str(p.get("series") or "Unknown")
        score = str(p["x"])
        if q_text not in matrix:
            matrix[q_text] = {"5": 0, "4": 0, "3": 0, "2": 0, "1": 0}
        matrix[q_text][score] = int(p["y"])

    cat_map = {c.id: c.name for c in survey.categories}
    rows = []
    for q in survey.questions:
        if q.text not in matrix:
            continue
        rows.append({
            "Category": cat_map.get(q.category_id, "Uncategorized"),
            "Question": q.text,
            "5": matrix[q.text]["5"],
            "4": matrix[q.text]["4"],
            "3": matrix[q.text]["3"],
            "2": matrix[q.text]["2"],
            "1": matrix[q.text]["1"],
        })

    df = pd.DataFrame(rows)
    if not df.empty:
        # Sort by category then question
        df = df.sort_values(["Category", "Question"])
    return df


CAT_COLORS = [
    "DDEBF7",  # Light Blue
    "E2EFDA",  # Light Green
    "FFF2CC",  # Light Yellow
    "FCE4D6",  # Light Orange
    "D9E1F2",  # Light Indigo
    "E7E6E6",  # Light Gray
    "F2F2F2",  # Lighter Gray
]


def _write_raw_responses_ws(ws, df: pd.DataFrame, col_metadata: list[dict[str, Any]]):
    """Specific writer for the Raw Responses sheet with category headers."""
    if df.empty:
        ws.cell(row=1, column=1, value="(no data)")
        return

    # 1. Prepare colors mapping
    unique_categories = []
    for m in col_metadata:
        cat = m["category"]
        if cat not in unique_categories:
            unique_categories.append(cat)
    
    cat_fills = {}
    for i, cat in enumerate(unique_categories):
        color = CAT_COLORS[i % len(CAT_COLORS)]
        cat_fills[cat] = PatternFill("solid", fgColor=color)

    # 2. Write Category Header (Row 1)
    for c, m in enumerate(col_metadata, start=1):
        cell = ws.cell(row=1, column=c, value=m["category"])
        cell.fill = cat_fills.get(m["category"], HEADER_FILL)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 3. Merge Category Cells
    if col_metadata:
        start_col = 1
        current_cat = col_metadata[0]["category"]
        for c, m in enumerate(col_metadata[1:], start=2):
            if m["category"] != current_cat:
                if (c - 1) > start_col:
                    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=c - 1)
                start_col = c
                current_cat = m["category"]
        # Final merge
        if len(col_metadata) > start_col:
            ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=len(col_metadata))

    # 4. Write Question Header (Row 2)
    for c, m in enumerate(col_metadata, start=1):
        cell = ws.cell(row=2, column=c, value=m["name"])
        cell.fill = cat_fills.get(m["category"], HEADER_FILL)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 5. Write Data (starting Row 3)
    for r_idx, (_, rec) in enumerate(df.iterrows(), start=3):
        for c_idx, m in enumerate(col_metadata, start=1):
            v = rec[m["id"]]
            if hasattr(v, "to_pydatetime"):
                v = v.to_pydatetime()
            ws.cell(row=r_idx, column=c_idx, value=None if pd.isna(v) else v)

    # 6. Auto-width columns and freeze panes
    ws.freeze_panes = "A3"
    for c, m in enumerate(col_metadata, start=1):
        col_letter = get_column_letter(c)
        if len(df):
            # Sample first few rows for width
            sample_vals = df[m["id"]].head(20).tolist()
            lens = [len(str(v)) for v in sample_vals if not pd.isna(v)]
            lens.append(len(str(m["name"])))
            max_len = max(lens) if lens else 12
        else:
            max_len = len(str(m["name"]))
        ws.column_dimensions[col_letter].width = max(10, min(50, max_len + 2))


# ---------- main entry ----------

def build_workbook(db: Session, req: schemas.ExportIn) -> bytes:
    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "Summary"

    surveys = db.scalars(
        select(models.Survey).where(models.Survey.id.in_(req.survey_ids))
    ).all()

    summary_rows = []
    for s in surveys:
        response_count = db.query(models.Response).filter(models.Response.survey_id == s.id).count()
        summary_rows.append(
            {
                "Survey": s.title,
                "Period": s.period_label,
                "State": s.state,
                "Created": s.created_at,
                "Opened": s.opened_at,
                "Closed": s.closed_at,
                "Responses": response_count,
            }
        )
    _write_df(ws_summary, pd.DataFrame(summary_rows), title="Surveys included")

    if req.include_aggregates:
        ws = wb.create_sheet("Aggregates")
        row = 1
        for s in surveys:
            row = _write_df(
                ws,
                aggregations.aggregate_per_question(db, s.id),
                start_row=row,
                title=f"[{s.title}] Per-question averages",
            )
            row = _write_df(
                ws,
                aggregations.category_avg_scores(db, s.id),
                start_row=row,
                title=f"[{s.title}] Per-category averages",
            )
            # Response counts by demographic (no avg score)
            for dim, label in [
                ("gender", "by gender"),
                ("occupation", "by occupation"),
                ("age_bucket", "by age bucket"),
            ]:
                row = _write_df(
                    ws,
                    aggregations.demographic_breakdown(db, s.id, dim),
                    start_row=row,
                    title=f"[{s.title}] Response count {label}",
                )
            # Per-question average by each demographic / category (rows=question, cols=value)
            for dim, label in [
                ("gender", "gender"),
                ("occupation", "occupation"),
                ("age_bucket", "age bucket"),
                ("category", "category"),
            ]:
                row = _write_df(
                    ws,
                    aggregations.category_x_question_avg(db, s.id, dim),
                    start_row=row,
                    title=f"[{s.title}] Avg score per question by {label}",
                )

    if req.include_raw:
        for s in surveys:
            name = f"Raw - {s.title}"[:31] or "Raw"
            ws = wb.create_sheet(name)
            df, meta = aggregations.raw_responses_df(db, s.id)
            _write_raw_responses_ws(ws, df, meta)

    if req.include_comparison and len(req.survey_ids) > 1:
        ws = wb.create_sheet("Comparison")
        _write_pivot(
            ws,
            aggregations.comparison_across_surveys(db, req.survey_ids),
            1,
            "Average score per question across selected surveys",
        )

    if req.include_occupation_breakdown:
        ws = wb.create_sheet("Occupation Breakdown")
        row = 1
        for s in surveys:
            # Get list of occupations for this survey
            occ_df = aggregations.demographic_breakdown(db, s.id, "occupation")
            if occ_df.empty:
                continue
            
            occupations = occ_df["occupation"].tolist()
            for occ in occupations:
                ws.cell(row=row, column=1, value=f"[{s.title}] Occupation: {occ}").font = Font(bold=True, size=13)
                row += 1

                # Run query for this occupation: metric=count, group_by=question, x_axis=score
                query = schemas.AnalyticsQueryIn(
                    survey_ids=[s.id],
                    metric="count",
                    x_axis="score",
                    group_by="question",
                    filters=schemas.AnalyticsFilter(occupation=[occ])
                )
                result = aggregations.run_analytics(db, query)
                points = [p.model_dump() for p in result.points]
                
                table = _occupation_breakdown_to_table(points, s)
                if table.empty:
                    ws.cell(row=row, column=1, value="(no data)")
                    row += 2
                    continue
                
                data_start_row = row
                row = _write_df(ws, table, start_row=row)
                data_rows = len(table)
                
                chart = _build_occupation_chart(f"{occ} - Score Distribution", data_rows, ws, data_start_row)
                anchor = f"{get_column_letter(10)}{data_start_row}"
                ws.add_chart(chart, anchor)
                row = max(row, data_start_row + 20)

    if req.charts:
        ws = wb.create_sheet("Charts")
        row = 1
        for ch in req.charts:
            # Title
            ws.cell(row=row, column=1, value=ch.name).font = Font(bold=True, size=13)
            row += 1

            cfg: dict[str, Any] = dict(ch.config_json)
            try:
                query = schemas.AnalyticsQueryIn(survey_ids=ch.survey_ids, **cfg)
                result = aggregations.run_analytics(db, query)
            except Exception as e:  # noqa: BLE001
                ws.cell(row=row, column=1, value=f"(query failed: {e})")
                row += 2
                continue

            points = [p.model_dump() for p in result.points]
            table = _analytics_points_to_table(points)
            if table.empty:
                ws.cell(row=row, column=1, value="(no data for this chart)")
                row += 2
                continue

            data_start_row = row  # header row of the data table
            row = _write_df(ws, table, start_row=row)
            data_rows = len(table)  # number of data rows (excluding header)
            data_cols = len(table.columns)

            chart_type = str(cfg.get("type", "bar"))
            chart = _build_native_chart(chart_type, ch.name, data_rows, data_cols, ws, data_start_row)
            anchor = f"{get_column_letter(data_cols + 2)}{data_start_row}"
            ws.add_chart(chart, anchor)
            # Leave vertical room for the chart image (roughly 18 rows)
            row = max(row, data_start_row + 20)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
